from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.db.models import Count, Q
from django.http import HttpResponse, JsonResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.utils import timezone
from django.utils.translation import gettext as _
from django.views.decorators.http import require_GET, require_POST
from openpyxl import Workbook
from datetime import datetime

from .forms import BatchForm, ClearRecordsForm, IngredientsForm, ProductForm
from .models import Batch, Product, TrackingLabel, IngredientLabel

# ---------------------------------------------------------------------------
# Main print page
# ---------------------------------------------------------------------------

@login_required
def print_page(request):
    """The main printing screen (search + the two independent print buttons)."""
    today = timezone.localdate()
    context = {
        "today_count": TrackingLabel.objects.filter(print_date=today).count(),
        "month_count": TrackingLabel.objects.filter(
            print_date__year=today.year, print_date__month=today.month
        ).count(),
    }
    return render(request, "labels/print.html", context)


@login_required
@require_GET
def api_search_products(request):
    """Live search used by the perfume-name search box."""
    q = request.GET.get("q", "").strip()
    products = Product.objects.filter(is_active=True)
    if q:
        products = products.filter(name__icontains=q)
    products = products[:10]
    data = [
        {
            "id": p.id,
            "name": p.name,
            "oil_batch_number": p.oil_batch_number or "",
            "batch": p.active_batch.batch_number if p.active_batch else "",
        }
        for p in products
    ]
    return JsonResponse({"results": data})


@login_required
@require_GET
def api_product_detail(request, product_id):
    product = get_object_or_404(Product, pk=product_id, is_active=True)
    today = timezone.localdate()
    
    # Generate a preview tracking number and store in session
    preview_number = _generate_preview_tracking_number(request, product)
    
    return JsonResponse(
        {
            "id": product.id,
            "name": product.name,
            "oil_batch_number": product.oil_batch_number or "",
            "batch": product.active_batch.batch_number if product.active_batch else "",
            "ingredients": product.ingredients_text,
            "preview_tracking": preview_number,
            "today_count": TrackingLabel.objects.filter(
                product=product, print_date=today
            ).count(),
            "month_count": TrackingLabel.objects.filter(
                product=product,
                print_date__year=today.year,
                print_date__month=today.month,
            ).count(),
        }
    )


def _generate_preview_tracking_number(request, product):
    """Generate and store a preview tracking number in session."""
    import random
    
    # Check if we already have a preview for this product in session
    session_key = f"preview_tracking_{product.id}"
    preview_number = request.session.get(session_key)
    
    if preview_number:
        # Check if it's still valid (not used yet)
        if not TrackingLabel.objects.filter(tracking_number=preview_number).exists():
            return preview_number
    
    # Generate a new unique tracking number
    date_code = timezone.localdate().strftime("%d%m%y")
    for _attempt in range(999):
        serial = f"{random.randint(1, 999):03d}"
        tracking_number = f"{date_code}-{serial}"
        if not TrackingLabel.objects.filter(tracking_number=tracking_number).exists():
            # Store in session
            request.session[session_key] = tracking_number
            return tracking_number
    
    return "------"


@login_required
@require_POST
def print_tracking(request):
    """Creates one permanent TrackingLabel row per bottle and returns the
    data needed to render/print the 15x30mm labels (printer #1)."""
    product_id = request.POST.get("product_id")
    quantity = max(1, min(100, int(request.POST.get("quantity", 1) or 1)))
    product = get_object_or_404(Product, pk=product_id, is_active=True)
    batch = product.active_batch
    if not batch:
        return JsonResponse(
            {"error": _("This perfume has no active batch yet.")}, status=400
        )

    requested_tracking_number = request.POST.get("tracking_number", "").strip().upper()
    if requested_tracking_number:
        existing_label = TrackingLabel.objects.filter(
            product=product, tracking_number=requested_tracking_number
        ).first()
        if existing_label is None:
            return JsonResponse(
                {"error": _("The selected tracking number was not found.")},
                status=404,
            )
        existing_label.reprint_count += quantity
        existing_label.save(update_fields=["reprint_count"])
        payload = [
            {
                "name": existing_label.product.name,
                "tracking_number": existing_label.tracking_number,
                "oil_batch_number": product.oil_batch_number or "",
                "batch": existing_label.batch.batch_number,
            }
            for _ in range(quantity)
        ]
        return JsonResponse({"printer": "tracking", "labels": payload})

    # Check if we have a preview tracking number in session
    session_key = f"preview_tracking_{product.id}"
    preview_number = request.session.pop(session_key, None)
    
    labels = []
    
    # First label: use preview number if available and valid
    if preview_number and not TrackingLabel.objects.filter(tracking_number=preview_number).exists():
        label = TrackingLabel.objects.create(
            product=product,
            batch=batch,
            product_name_snapshot=product.name,
            batch_number_snapshot=batch.batch_number,
            print_date=timezone.localdate(),
            serial=preview_number.split('-')[1] if '-' in preview_number else '---',
            tracking_number=preview_number,
            printed_by=request.user,
        )
        labels.append(label)
        quantity -= 1
    
    # Generate remaining labels
    for _i in range(quantity):
        label = TrackingLabel.generate_unique(product, batch, user=request.user)
        labels.append(label)
    
    payload = [
        {
            "name": l.product.name,
            "tracking_number": l.tracking_number,
            "oil_batch_number": product.oil_batch_number or "",
            "batch": l.batch.batch_number,
        }
        for l in labels
    ]
    return JsonResponse({"printer": "tracking", "labels": payload})


@login_required
@require_POST
def print_ingredients(request):
    """Renders the 56x56mm ingredients label (printer #2). Saves to database."""
    product_id = request.POST.get("product_id")
    quantity = max(1, min(100, int(request.POST.get("quantity", 1) or 1)))
    note = (request.POST.get("note") or "").strip()
    product = get_object_or_404(Product, pk=product_id, is_active=True)
    if not product.ingredients_text:
        return JsonResponse(
            {"error": _("Add this perfume's ingredients first.")}, status=400
        )
    
    # Save ingredient label to database
    IngredientLabel.objects.create(
        product=product,
        product_name_snapshot=product.name,
        oil_batch_snapshot=product.oil_batch_number,
        print_date=timezone.localdate(),
        quantity=quantity,
        note=note,
        printed_by=request.user,
    )
    
    return JsonResponse(
        {
            "printer": "ingredients",
            "name": product.name,
            "ingredients": product.ingredients_text,
            "note": note,
            "quantity": quantity,
        }
    )


@login_required
@require_POST
def reprint_tracking(request):
    """Reprints an EXISTING tracking label without ever generating a new
    tracking number (per spec)."""
    tracking_number = request.POST.get("tracking_number", "").strip().upper()
    label = TrackingLabel.objects.filter(tracking_number=tracking_number).first()
    if not label:
        return JsonResponse({"error": _("Tracking number not found.")}, status=404)
    label.reprint_count += 1
    label.save(update_fields=["reprint_count"])
    return JsonResponse(
        {
            "printer": "tracking",
            "labels": [
                {
                    "name": label.product.name,
                    "tracking_number": label.tracking_number,
                    "oil_batch_number": label.product.oil_batch_number or "",
                    "batch": label.batch.batch_number,
                }
            ],
        }
    )


# ---------------------------------------------------------------------------
# Product & batch management
# ---------------------------------------------------------------------------

@login_required
def products_page(request):
    edit_id = request.GET.get("edit")
    editing = get_object_or_404(Product, pk=edit_id) if edit_id else None

    if request.method == "POST":
        instance = editing if "product_id" in request.POST else None
        is_new_product = instance is None
        form = ProductForm(request.POST, instance=instance)
        if form.is_valid():
            product = form.save()

            if is_new_product:
                # Create a default batch for the product
                batch_number = Batch.generate_unique_number()
                Batch.objects.create(
                    product=product,
                    batch_number=batch_number,
                    is_active=True,
                )

            messages.success(request, _("Product saved."))
            return redirect("labels:products")
    else:
        form = ProductForm(instance=editing)

    batch_form = BatchForm()
    context = {
        "form": form,
        "batch_form": batch_form,
        "editing": editing,
        "products": Product.objects.all().prefetch_related("batches"),
    }
    return render(request, "labels/products.html", context)


@login_required
@require_POST
def product_add_batch(request, product_id):
    product = get_object_or_404(Product, pk=product_id)
    form = BatchForm(request.POST)
    if form.is_valid():
        batch = form.save(commit=False)
        batch.product = product
        batch.save()
        messages.success(request, _("Batch added."))
    else:
        messages.error(request, _("Couldn't add that batch — please check the batch number."))
    return redirect("labels:products")


@login_required
@require_POST
def product_set_active_batch(request, product_id, batch_id):
    batch = get_object_or_404(Batch, pk=batch_id, product_id=product_id)
    batch.is_active = True
    batch.save()
    messages.success(request, _("Active batch updated."))
    return redirect("labels:products")


from django.db.models import ProtectedError

@login_required
@require_POST
def product_delete(request, product_id):
    """Permanently deletes a perfume. Its Batches are deleted too (CASCADE),
    but existing TrackingLabel rows survive with product/batch set to null —
    their snapshot fields keep the name/batch readable."""
    product = get_object_or_404(Product, pk=product_id)
    name = product.name
    product.delete()
    messages.success(request, _("Perfume '%(name)s' deleted.") % {"name": name})
    return redirect("labels:products")

# ---------------------------------------------------------------------------
# Ingredients management
# ---------------------------------------------------------------------------


@login_required
def ingredients_page(request):
    product_id = request.POST.get("product_id") or request.GET.get("product")
    product = None
    if product_id:
        product = get_object_or_404(Product, pk=product_id)

    if request.method == "POST" and product:
        form = IngredientsForm(request.POST, instance=product)
        if form.is_valid():
            form.save()
            messages.success(request, _("Ingredients saved for %(name)s") % {"name": product.name})
            return redirect(f"{request.path}?product={product.id}")
    else:
        form = IngredientsForm(instance=product)

    context = {
        "form": form,
        "product": product,
        "products": Product.objects.all(),
    }
    return render(request, "labels/ingredients.html", context)


# ---------------------------------------------------------------------------
# Reports & search
# ---------------------------------------------------------------------------


@login_required
def reports_page(request):
    query = request.GET.get("q", "").strip()
    
    # Get tracking labels
    tracking_labels = TrackingLabel.objects.select_related("product", "batch").all()
    
    # Get ingredient labels
    ingredient_labels = IngredientLabel.objects.select_related("product").all()
    
    # Combine and filter
    all_records = []
    
    # Add tracking labels
    for label in tracking_labels:
        all_records.append({
            'type': 'tracking',
            'created_at': label.created_at,
            'print_date': label.print_date,
            'product_name': label.product.name if label.product else label.product_name_snapshot,
            'number': label.tracking_number,
            'oil_batch': label.product.oil_batch_number if label.product else '',
            'quantity': 1,
            'note': '',
            'label_object': label,
        })
    
    # Add ingredient labels
    for label in ingredient_labels:
        all_records.append({
            'type': 'ingredient',
            'created_at': label.created_at,
            'print_date': label.print_date,
            'product_name': label.product_name_snapshot,
            'number': label.oil_batch_snapshot,
            'oil_batch': label.oil_batch_snapshot,
            'quantity': label.quantity,
            'note': label.note,
            'label_object': label,
        })
    
    # Sort by created_at descending (newest first)
    all_records.sort(key=lambda x: x['created_at'], reverse=True)
    
    # Apply search filter if query exists
    if query:
        query_lower = query.lower()
        all_records = [
            r for r in all_records
            if query_lower in r['product_name'].lower()
            or query_lower in r['number'].lower()
            or query_lower in r['oil_batch'].lower()
        ]
    
    # Limit to 200 records
    all_records = all_records[:200]

    today = timezone.localdate()
    counters = []
    for product in Product.objects.all().prefetch_related("batches"):
        product_labels = TrackingLabel.objects.filter(product=product)
        counters.append({
            "product": product,
            "today": product_labels.filter(print_date=today).count(),
            "month": product_labels.filter(
                print_date__year=today.year, print_date__month=today.month
            ).count(),
            "total": product_labels.count(),
        })

    context = {
        "query": query,
        "records": all_records,
        "counters": counters,
        "today_count": TrackingLabel.objects.filter(print_date=today).count(),
        "month_count": TrackingLabel.objects.filter(
            print_date__year=today.year, print_date__month=today.month
        ).count(),
        "ingredients_today": IngredientLabel.objects.filter(print_date=today).count(),
        "products": Product.objects.all(),
    }
    return render(request, "labels/reports.html", context)


def models_q_search(query):
    from django.db.models import Q

    return (
        Q(tracking_number__icontains=query)
        | Q(product__name__icontains=query)
        | Q(batch__batch_number__icontains=query)
    )


@login_required
@require_POST
def clear_tracking_records(request):
    """Permanently deletes tracking-label records printed within a chosen
    date range (inclusive). Requires an explicit start and end date —
    there is no "clear everything" shortcut, to avoid accidents."""
    form = ClearRecordsForm(request.POST)
    if not form.is_valid():
        messages.error(request, _("Choose a valid start and end date."))
        return redirect("labels:reports")

    start_date = form.cleaned_data["start_date"]
    end_date = form.cleaned_data["end_date"]
    qs = TrackingLabel.objects.filter(print_date__gte=start_date, print_date__lte=end_date)
    count = qs.count()
    qs.delete()

    if count:
        messages.success(
            request,
            _("Deleted %(count)s tracking record(s) from %(start)s to %(end)s.")
            % {"count": count, "start": start_date.isoformat(), "end": end_date.isoformat()},
        )
    else:
        messages.info(
            request,
            _("No tracking records found between %(start)s and %(end)s.")
            % {"start": start_date.isoformat(), "end": end_date.isoformat()},
        )
    return redirect("labels:reports")


# Excel export column headers are always German
_XLSX_HEADERS = ["Parfümname", "Trackingnummer", "Chargennummer", "Inhaltsstoffe"]
_XLSX_TOP_SELLERS_HEADER = "Top 50 meistverkaufte Parfums"
_XLSX_QUANTITY_HEADER = "Verkaufte Menge"


def _xlsx_response(workbook, filename):
    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    workbook.save(response)
    return response


def _filtered_labels(request):
    query = request.GET.get("q", "").strip()
    labels = TrackingLabel.objects.select_related("product", "batch").order_by("id")
    if query:
        labels = labels.filter(models_q_search(query))
    return labels


def _write_records_sheet(ws, labels):
    ws.append(_XLSX_HEADERS)
    for label in labels:
        ws.append(
            [
                label.product.name,
                label.tracking_number,
                label.batch.batch_number,
                label.product.ingredients_text,
            ]
        )
    ws.column_dimensions["E"].width = 10


@login_required
def export_xlsx(request):
    """Excel export: one row per printed bottle."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    _write_records_sheet(ws, _filtered_labels(request))

    filename = f"Berichte-{timezone.localdate().isoformat()}.xlsx"
    return _xlsx_response(wb, filename)


@login_required
def export_xlsx_with_counter(request):
    """Same export as above, plus a Top-50 best-selling-perfumes summary."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    _write_records_sheet(ws, _filtered_labels(request))

    ws["F1"] = _XLSX_TOP_SELLERS_HEADER
    ws["G1"] = _XLSX_QUANTITY_HEADER
    top_sellers = (
        TrackingLabel.objects.values("product__name")
        .annotate(quantity=Count("id"))
        .order_by("-quantity")[:50]
    )
    for row_index, row in enumerate(top_sellers, start=2):
        ws.cell(row=row_index, column=6, value=row["product__name"])
        ws.cell(row=row_index, column=7, value=row["quantity"])

    filename = f"Berichte-Zähler-{timezone.localdate().isoformat()}.xlsx"
    return _xlsx_response(wb, filename)